"""
Generate Internship Offer Letter DOCX with Horrazon color theme.

Color Theme (from email template):
- Header gradient: #0f2544 → #1e3a6e → #2563eb
- Primary blue accent: #1e40af
- Dark text: #1e293b
- Body text: #374151
- Light background: #f8fafc
- Border/subtle: #e2e8f0
- Muted text: #64748b
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

# ─── Color Constants (Horrazon Theme) ───────────────────────────────────────
NAVY_DARK = RGBColor(0x0F, 0x25, 0x44)  # #0f2544
NAVY_MID = RGBColor(0x1E, 0x3A, 0x6E)  # #1e3a6e
BLUE_PRIMARY = RGBColor(0x1E, 0x40, 0xAF)  # #1e40af
BLUE_BRIGHT = RGBColor(0x25, 0x63, 0xEB)  # #2563eb
DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)  # #1e293b
BODY_TEXT = RGBColor(0x37, 0x41, 0x51)  # #374151
MUTED_TEXT = RGBColor(0x64, 0x74, 0x8B)  # #64748b
WHITE = RGBColor(0xFF, 0xFF, 0xFF)  # #ffffff
LIGHT_BG = "F8FAFC"  # for table shading
BORDER_COLOR = "E2E8F0"  # #e2e8f0


def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_border(cell, **kwargs):
    """Set cell borders. kwargs: top, bottom, left, right with value (size, color, style)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')

    for edge, attrs in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{attrs.get("style", "single")}" '
            f'w:sz="{attrs.get("size", "4")}" w:space="0" '
            f'w:color="{attrs.get("color", "000000")}"/>'
        )
        tcBorders.append(element)

    tcPr.append(tcBorders)


def add_horizontal_line(doc, color="1E40AF"):
    """Add a colored horizontal line."""
    paragraph = doc.add_paragraph()
    paragraph.space_after = Pt(6)
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="{color}"/>'
        f"</w:pBdr>"
    )
    pPr.append(pBdr)


def create_offer_letter():
    doc = Document()

    # ─── Page Setup ──────────────────────────────────────────────────────────
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # ─── Header Banner ───────────────────────────────────────────────────────
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cell = header_table.cell(0, 0)
    set_cell_shading(header_cell, "0F2544")

    # Company Name in Header
    p = header_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(20)
    p.space_after = Pt(4)
    run = p.add_run("HORRAZON INTELLIGENCE PRIVATE LIMITED")
    run.font.size = Pt(18)
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.name = "Segoe UI"

    # Tagline
    p2 = header_cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.space_after = Pt(20)
    run2 = p2.add_run("AI-Powered Business Intelligence")
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0xA0, 0xB8, 0xD8)
    run2.font.name = "Segoe UI"

    doc.add_paragraph()  # spacer

    # ─── Title ───────────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.space_after = Pt(4)
    run = title.add_run("INTERNSHIP OFFER LETTER")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = NAVY_DARK
    run.font.name = "Segoe UI"

    add_horizontal_line(doc, "1E40AF")

    # ─── Reference & Date ────────────────────────────────────────────────────
    ref_para = doc.add_paragraph()
    ref_para.space_after = Pt(2)
    run = ref_para.add_run("Ref. No.: ")
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED_TEXT
    run.font.name = "Segoe UI"
    run = ref_para.add_run("HZG/2026/AIML-INT/02")
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_TEXT
    run.font.bold = True
    run.font.name = "Segoe UI"

    date_para = doc.add_paragraph()
    date_para.space_after = Pt(16)
    run = date_para.add_run("Date: ")
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED_TEXT
    run.font.name = "Segoe UI"
    run = date_para.add_run("03 August 2026")
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_TEXT
    run.font.bold = True
    run.font.name = "Segoe UI"

    # ─── To Section ──────────────────────────────────────────────────────────
    to_para = doc.add_paragraph()
    to_para.space_after = Pt(2)
    run = to_para.add_run("To,")
    run.font.size = Pt(11)
    run.font.color.rgb = BODY_TEXT
    run.font.name = "Segoe UI"

    lines = [
        ("Ms. Tanishka Chauhan", True),
        ("Bachelor of Technology (Artificial Intelligence)", False),
        ("Graphic Era Deemed to be University", False),
        ("Batch: 2022–2026", False),
    ]
    for text, bold in lines:
        p = doc.add_paragraph()
        p.space_after = Pt(1)
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_TEXT if bold else BODY_TEXT
        run.font.bold = bold
        run.font.name = "Segoe UI"

    doc.add_paragraph()  # spacer

    # ─── Subject ─────────────────────────────────────────────────────────────
    subject_table = doc.add_table(rows=1, cols=1)
    subject_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    subject_cell = subject_table.cell(0, 0)
    set_cell_shading(subject_cell, LIGHT_BG)
    set_cell_border(
        subject_cell, left={"size": "18", "color": "1E40AF", "style": "single"}
    )

    p = subject_cell.paragraphs[0]
    p.space_before = Pt(8)
    p.space_after = Pt(8)
    run = p.add_run("Subject: ")
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED_TEXT
    run.font.name = "Segoe UI"
    run = p.add_run("Offer of Internship – AI Engineer Intern (LLM & Agentic AI)")
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE_PRIMARY
    run.font.bold = True
    run.font.name = "Segoe UI"

    doc.add_paragraph()  # spacer

    # ─── Dear... ─────────────────────────────────────────────────────────────
    def add_body_text(text, bold=False, space_after=8, color=None):
        p = doc.add_paragraph()
        p.space_after = Pt(space_after)
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.color.rgb = color or BODY_TEXT
        run.font.bold = bold
        run.font.name = "Segoe UI"
        return p

    add_body_text("Dear Ms. Tanishka Chauhan,", color=DARK_TEXT)

    add_body_text(
        "Following the successful completion of our recruitment and technical evaluation process, "
        "we are pleased to offer you the position of AI Engineer Intern (LLM & Agentic AI) at "
        'Horrazon Intelligence Private Limited ("the Company").'
    )

    # Congratulations
    congrats = doc.add_paragraph()
    congrats.space_before = Pt(8)
    congrats.space_after = Pt(8)
    run = congrats.add_run("Congratulations on your selection.")
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = BLUE_PRIMARY
    run.font.name = "Segoe UI"

    add_body_text(
        "Throughout our evaluation process, you demonstrated strong technical aptitude, "
        "sound engineering judgment, problem-solving ability, and a genuine commitment to "
        "learning and building AI-powered systems. We believe your skills and potential align "
        "well with the engineering culture and long-term vision of Horrazon AI, and we look "
        "forward to welcoming you to our team."
    )

    add_body_text(
        "This Internship Offer, together with Schedule A – Terms & Conditions of Internship "
        "Engagement, constitutes the complete terms governing your internship engagement with the Company."
    )

    doc.add_paragraph()  # spacer

    # ─── Internship Details Table ────────────────────────────────────────────
    details_heading = doc.add_paragraph()
    details_heading.space_after = Pt(8)
    run = details_heading.add_run("INTERNSHIP DETAILS")
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = NAVY_DARK
    run.font.name = "Segoe UI"

    details_data = [
        ("Designation", "AI Engineer Intern (LLM & Agentic AI)"),
        ("Engagement Type", "Fixed-Term Internship"),
        ("Work Mode", "Remote"),
        ("Joining Date", "04 August 2026"),
        ("Internship Duration", "Two (2) Months*"),
        ("Reporting Manager", "Co-Founder & Chief Technology Officer"),
        ("Department", "Engineering & AI"),
        ("Monthly Stipend", "₹10,000 (Rupees Ten Thousand Only)"),
    ]

    table = doc.add_table(rows=len(details_data) + 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(["Particular", "Details"]):
        set_cell_shading(hdr_cells[i], "1E40AF")
        p = hdr_cells[i].paragraphs[0]
        p.space_before = Pt(6)
        p.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = WHITE
        run.font.name = "Segoe UI"

    # Data rows
    for idx, (key, value) in enumerate(details_data):
        row_cells = table.rows[idx + 1].cells
        bg_color = LIGHT_BG if idx % 2 == 0 else "FFFFFF"

        # Key cell
        set_cell_shading(row_cells[0], bg_color)
        p = row_cells[0].paragraphs[0]
        p.space_before = Pt(5)
        p.space_after = Pt(5)
        run = p.add_run(key)
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = DARK_TEXT
        run.font.name = "Segoe UI"

        # Value cell
        set_cell_shading(row_cells[1], bg_color)
        p = row_cells[1].paragraphs[0]
        p.space_before = Pt(5)
        p.space_after = Pt(5)
        run = p.add_run(value)
        run.font.size = Pt(10)
        run.font.color.rgb = BODY_TEXT
        run.font.name = "Segoe UI"

    # Set column widths
    for row in table.rows:
        row.cells[0].width = Cm(5.5)
        row.cells[1].width = Cm(11.5)

    # Footnote
    footnote = doc.add_paragraph()
    footnote.space_before = Pt(4)
    footnote.space_after = Pt(16)
    run = footnote.add_run(
        "*Subject to extension based on performance, project requirements, and business needs, "
        "at the sole discretion of the Company."
    )
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = MUTED_TEXT
    run.font.name = "Segoe UI"

    # ─── Acceptance of Offer ─────────────────────────────────────────────────
    doc.add_paragraph()

    acceptance_heading = doc.add_paragraph()
    acceptance_heading.space_after = Pt(8)
    run = acceptance_heading.add_run("ACCEPTANCE OF OFFER")
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = NAVY_DARK
    run.font.name = "Segoe UI"

    add_horizontal_line(doc, "2563EB")

    add_body_text(
        "This offer shall remain valid for 48 hours from the date of issuance unless "
        "otherwise extended in writing by the Company."
    )

    add_body_text(
        "Acceptance of this offer shall be completed through the Zoho Sign workflow "
        "within the stipulated acceptance period."
    )

    add_body_text(
        "Upon receipt of the completed documents, the Company will initiate your onboarding "
        "process and share your joining instructions and other relevant onboarding information."
    )

    doc.add_paragraph()

    # ─── Closing ─────────────────────────────────────────────────────────────
    # Motivational box
    motiv_table = doc.add_table(rows=1, cols=1)
    motiv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    motiv_cell = motiv_table.cell(0, 0)
    set_cell_shading(motiv_cell, "EEF2FF")

    p = motiv_cell.paragraphs[0]
    p.space_before = Pt(12)
    p.space_after = Pt(12)
    run = p.add_run(
        "We are pleased to welcome you to Horrazon AI and look forward to your contributions "
        "as part of our Engineering team. We wish you every success in your internship and "
        "look forward to working with you."
    )
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE_PRIMARY
    run.font.italic = True
    run.font.name = "Segoe UI"

    doc.add_paragraph()

    # ─── Signature ───────────────────────────────────────────────────────────
    add_body_text("For Horrazon Intelligence Private Limited", color=BODY_TEXT)
    doc.add_paragraph()  # space for signature
    doc.add_paragraph()  # space for signature

    sig_name = doc.add_paragraph()
    sig_name.space_after = Pt(2)
    run = sig_name.add_run("Karan Singh")
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Segoe UI"

    sig_title = doc.add_paragraph()
    sig_title.space_after = Pt(2)
    run = sig_title.add_run("Co-Founder & Chief Technology Officer")
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED_TEXT
    run.font.name = "Segoe UI"

    sig_company = doc.add_paragraph()
    sig_company.space_after = Pt(2)
    run = sig_company.add_run("Horrazon Intelligence Private Limited")
    run.font.size = Pt(10)
    run.font.color.rgb = BLUE_PRIMARY
    run.font.bold = True
    run.font.name = "Segoe UI"

    doc.add_paragraph()
    add_horizontal_line(doc, "E2E8F0")
    doc.add_paragraph()

    # ─── Candidate Acknowledgement ──────────────────────────────────────────
    ack_heading = doc.add_paragraph()
    ack_heading.space_after = Pt(8)
    run = ack_heading.add_run("CANDIDATE ACKNOWLEDGEMENT")
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = NAVY_DARK
    run.font.name = "Segoe UI"

    add_horizontal_line(doc, "2563EB")

    add_body_text(
        "I acknowledge that I have read, understood, and accepted this Internship Offer "
        "together with Schedule A – Terms & Conditions of Internship Engagement.",
        space_after=20,
    )

    # Signature fields
    fields = [
        "Candidate Name: __________________________________________",
        "Signature: _______________________________________________",
        "Date: ____________________________________________________",
    ]
    for field in fields:
        p = doc.add_paragraph()
        p.space_after = Pt(16)
        run = p.add_run(field)
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_TEXT
        run.font.name = "Segoe UI"

    # ─── Footer ──────────────────────────────────────────────────────────────
    doc.add_paragraph()

    footer_table = doc.add_table(rows=1, cols=1)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer_cell = footer_table.cell(0, 0)
    set_cell_shading(footer_cell, "0F2544")

    p = footer_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(12)
    p.space_after = Pt(12)
    run = p.add_run(
        "Horrazon Intelligence Pvt. Ltd. · Empowering Enterprises Through AI"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    run.font.name = "Segoe UI"

    # ─── Save ────────────────────────────────────────────────────────────────
    output_path = os.path.join(
        os.path.dirname(__file__), "Internship_Offer_Letter_Tanishka_Chauhan.docx"
    )
    doc.save(output_path)
    print(f"✅ Offer letter generated: {output_path}")
    return output_path


if __name__ == "__main__":
    create_offer_letter()
